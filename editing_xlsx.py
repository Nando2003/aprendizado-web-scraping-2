from openpyxl import load_workbook
from openpyxl.styles import PatternFill

def editing_xlsx(excel_path:str, data:list, linha:int =1, coluna:str ='A') -> bool:
    """ data = [(data, color), (data,color)...] """
    try:
        load_wb = load_workbook(excel_path)
        load_ws = load_wb.active
        
        for i, (value, color) in enumerate(data, start=linha):
            cell = load_ws[f'{coluna}{i}']
            cell.value = value
            if color:
                cell.fill = PatternFill(
                    start_color=color,
                    end_color=color,
                    fill_type='solid'
                )
                
        load_wb.save(excel_path)
        return True
    
    except Exception:
        return False